"""Table I/O that preserves identity text and records private provenance."""

from __future__ import annotations

import csv
import datetime
import hashlib
import json
import re
from pathlib import Path

import pandas as pd

from .security import SafeError, reject_links


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_json(path: Path) -> dict:
    def unique(pairs):
        result = {}
        for k, v in pairs:
            if k in result:
                raise SafeError("DUPLICATE_CONFIGURATION_KEY")
            result[k] = v
        return result
    reject_links(path)
    try:
        result = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=unique,
                            parse_constant=lambda _: (_ for _ in ()).throw(SafeError("NONFINITE_JSON")))
    except (OSError, ValueError):
        raise SafeError("INVALID_JSON_FILE") from None
    if not isinstance(result, dict):
        raise SafeError("CONFIGURATION_OBJECT_REQUIRED")
    return result


def write_json(path: Path, value) -> None:
    path.write_text(json.dumps(value, indent=2, sort_keys=True, allow_nan=False) + "\n",
                    encoding="utf-8")


def read_table(path: Path) -> pd.DataFrame:
    reject_links(path)
    # An LFS pointer is not input data. Fail instead of treating it as a table.
    with path.open("rb") as f:
        if f.read(128).startswith(b"version https://git-lfs.github.com/spec/" + b"v1"):
            raise SafeError("INPUT_NOT_MATERIALIZED")
    if path.suffix.lower() in {".csv", ".tsv"}:
        sep = "\t" if path.suffix.lower() == ".tsv" else ","
        with path.open(encoding="utf-8-sig", newline="") as f:
            header = next(csv.reader(f, delimiter=sep), [])
        frame = pd.read_csv(path, sep=sep, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = book.worksheets[0].iter_rows(values_only=True)
            header = [str(v) if v is not None else "" for v in next(rows, ())]
            def cell(value):
                if isinstance(value, datetime.datetime):
                    return value.date().isoformat()
                if isinstance(value, datetime.date):
                    return value.isoformat()
                return "" if value is None else str(value)
            frame = pd.DataFrame(([cell(v) for v in row] for row in rows), columns=header)
        finally:
            book.close()
    else:
        raise SafeError("UNSUPPORTED_INPUT_FORMAT")
    header = [h.strip() for h in header]
    if not header or any(not h for h in header) or len(set(header)) != len(header):
        raise SafeError("INVALID_OR_DUPLICATE_HEADERS")
    if any(h.startswith("_") for h in header):
        raise SafeError("RESERVED_INPUT_HEADER")
    frame.columns = header
    return frame


def load_sources(root: Path, files: list[str], columns: dict[str, str]) -> tuple[pd.DataFrame, list[dict]]:
    required = {"section", "date", "pci", "surface", "group"}
    if not required <= columns.keys() or set(columns) - required - {"construction"}:
        raise SafeError("INVALID_COLUMN_MAPPING")
    if len(set(columns.values())) != len(columns):
        raise SafeError("DUPLICATE_COLUMN_MAPPING")
    frames, manifest, seen = [], [], set()
    for relative in files:
        path = root / relative
        reject_links(path)
        path = path.resolve()
        if root.resolve() not in path.parents or path in seen:
            raise SafeError("INVALID_OR_DUPLICATE_INPUT_PATH")
        seen.add(path)
        before = sha256(path)
        frame = read_table(path)
        if not set(columns.values()) <= set(frame.columns):
            raise SafeError("REQUIRED_INPUT_COLUMNS_MISSING")
        if sha256(path) != before:
            raise SafeError("INPUT_CHANGED_DURING_READ")
        frame["_source_file"] = str(path.relative_to(root))
        frame["_source_row"] = range(2, len(frame) + 2)
        frame["_source_sha256"] = before
        # Preserve structural column absence across heterogeneous source exports.
        deduct_schema = tuple(c for c in frame if re.fullmatch(r"\d{2}[LMH]?\s+deduct", c, re.I))
        frame["_deduct_schema"] = [deduct_schema] * len(frame)
        frames.append(frame)
        manifest.append({"path": str(path.relative_to(root)), "sha256": before, "rows": len(frame)})
    if not frames:
        raise SafeError("NO_INPUT_FILES")
    return pd.concat(frames, ignore_index=True), manifest


def apply_corrections(frame: pd.DataFrame, path: Path | None) -> tuple[pd.DataFrame, pd.DataFrame]:
    cols = ["correction_id", "source_file", "source_row", "column", "expected", "replacement",
            "reason", "reference", "author", "date"]
    if path is None:
        return frame, pd.DataFrame(columns=cols)
    changes = read_table(path)
    if set(changes.columns) != set(cols) or changes.correction_id.duplicated().any():
        raise SafeError("INVALID_CORRECTIONS_TABLE")
    if changes[["correction_id", "reason", "reference", "author", "date"]].apply(
        lambda s: s.str.strip().eq("").any()).any():
        raise SafeError("CORRECTION_PROVENANCE_REQUIRED")
    result = frame.copy()
    for item in changes.to_dict("records"):
        if item["column"].startswith("_") or item["column"] not in result:
            raise SafeError("CORRECTION_COLUMN_INVALID")
        try:
            row = int(item["source_row"])
        except ValueError:
            raise SafeError("CORRECTION_ROW_INVALID") from None
        mask = result._source_file.eq(item["source_file"]) & result._source_row.eq(row)
        if mask.sum() != 1 or str(result.loc[mask, item["column"]].iloc[0]) != item["expected"]:
            raise SafeError("CORRECTION_EXPECTATION_MISMATCH")
        result.loc[mask, item["column"]] = item["replacement"]
    return result, changes
